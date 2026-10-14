#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
인게임 메일 텍스트 (mail_produce.mpc / mail_viewer.mpc) — dump / export-xlsx

구조 (2026-10-15 실측, PC 포트 복호화본 기준 — PS3판과 레이아웃 동일 계열):
  mail_produce.mpc = pac(엔트리 10, dance pac과 동일 레이아웃: +0x20 count,
  0x40 테이블 32B(+8 size,+12 ofs), MSG 매니페스트(파일명표), 데이터 연속)
    ent sya_mail_01.xmb  : 아이돌 메일 69 _text (id_sub/id_body/id_atc + n_*)
    ent sys_mail_01.xmb  : 샵/시스템 메일 42 _text (pack_* + n_*)
    ent all_getmail_01.scb(dec 래퍼 pac: SCB 1,472B + 컷씬 pac 64B)
       → 메일 수신 연출 대사 3행 — **gscript 번역 체계가 이미 커버**
         (웹앱 script/all의 all_getmail_01.scb 행). 내장 사본 리빌드는
         patch_attract.py 패턴(§20) 적용 대상. 116B 캡도 그 경로에서 처리.
    ent key_help.xmb     : 메일 화면 키 도움말 60 _text
  mail_viewer.mpc = pac(엔트리 10)
    sya_mail_01.xmb / key_help.xmb : produce와 **텍스트 완전 동일** (한 번
    번역해 두 파일에 주입), + viewer 전용 jou XMB(기자명 5 _text)

번역 행: 유니크 176 (sya 69 + sys 42 + key_help 60 + jou 5).
XMB 본문 최장 1,728B(UTF-16) — 가사와 같은 XMB 경로라 MSG 116B 캡 비적용.
한국어가 길어지면 rebuild_xmb(allow_grow)+pac 재조립(dance_lyrics.rebuild_pac)로
엔트리를 키운다(다음 단계 build에서).

  python mail_text.py dump         # mail_texts.json + xmb_translations.json에 mail/* 틀 합류
  python mail_text.py export-xlsx  # 게임 tools에 mail_translation_1.xlsx (기존 파일 보호)
"""
import os
import sys
import json
import argparse

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from xmb_text import parse_xmb, collect_texts  # noqa: E402
from dance_lyrics import decrypt_cached  # noqa: E402
from extract_image_text import JP_ROOT  # noqa: E402

XLSX_PATH = os.path.join(JP_ROOT, "gscript", "tools", "mail_translation_1.xlsx")

# (키, 라벨, 소스 mpc, 'XMB ' 절대 오프셋 탐색 식별자)
SOURCES = [
    ("mail/sya_mail_01.xmb", "아이돌 메일 (제목/본문/첨부)", "mail_produce.mpc", "id_sub"),
    ("mail/sys_mail_01.xmb", "샵·시스템 메일", "mail_produce.mpc", "pack_sub"),
    ("mail/key_help.xmb", "메일 화면 키 도움말", "mail_produce.mpc", "hel_000"),
    ("mail/jou_names.xmb", "기자 이름 (viewer)", "mail_viewer.mpc", "jou_000"),
]
# 주입 대상: 키 → 포함하는 mpc 목록 (viewer의 sya/key_help는 produce와 동일 텍스트)
INJECT_INTO = {
    "mail/sya_mail_01.xmb": ["mail_produce.mpc", "mail_viewer.mpc"],
    "mail/sys_mail_01.xmb": ["mail_produce.mpc"],
    "mail/key_help.xmb": ["mail_produce.mpc", "mail_viewer.mpc"],
    "mail/jou_names.xmb": ["mail_viewer.mpc"],
}


def find_xmb_with(data: bytes, marker: str):
    """marker id를 가진 XMB의 절대 오프셋."""
    p = 0
    while True:
        i = data.find(b"XMB ", p)
        if i < 0:
            raise ValueError(f"XMB({marker}) 없음")
        try:
            px = parse_xmb(data, i)
            ts = collect_texts(px)
            if any(t["id"] == marker for t in ts):
                return i, ts
        except Exception:  # noqa: BLE001 — 후보 스킵
            pass
        p = i + 4


def cmd_dump(_args):
    cache = {}
    table = {}
    locs = {}
    for key, label, mpc, marker in SOURCES:
        if mpc not in cache:
            cache[mpc] = decrypt_cached(os.path.join("mail", mpc))
        off, ts = find_xmb_with(cache[mpc], marker)
        rows = [{"id": t["id"], "jp": t["text"]} for t in ts]
        table[key] = rows
        locs[key] = {"label": label, "src": mpc, "xmb_off": off,
                     "inject_into": INJECT_INTO[key]}
        nz = sum(1 for r in rows if r["jp"])
        mx = max((len(r["jp"].encode("utf-16-be")) for r in rows), default=0)
        print(f"{key}: {len(rows)}행 (내용 {nz}) 최장 {mx}B — {label} [@{off:#x} in {mpc}]")
    out = {"locs": locs, "texts": table}
    with open(os.path.join(HERE, "mail_texts.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    # 번역 틀 합류 (기존 kr 보존)
    tr_path = os.path.join(HERE, "xmb_translations.json")
    tr = json.load(open(tr_path, encoding="utf-8"))
    n_new = 0
    for key, rows in table.items():
        ent = tr.setdefault(key, {})
        for r in rows:
            if r["id"] not in ent:
                ent[r["id"]] = {"jp": r["jp"], "cn": "", "kr": ""}
                n_new += 1
            else:
                ent[r["id"]]["jp"] = r["jp"]
    with open(tr_path, "w", encoding="utf-8") as f:
        json.dump(tr, f, ensure_ascii=False, indent=1)
    total = sum(len(r) for r in table.values())
    print(f"\n합계 {total}행 (신규 틀 {n_new}) → mail_texts.json + xmb_translations.json(mail/*)")
    print("참고: 메일 수신 연출 대사(all_getmail_01.scb 3행)는 기존 gscript 체계가 커버 —")
    print("      내장 SCB 사본 리빌드는 patch_attract 패턴으로 build 단계에서 처리 예정")


def cmd_export(_args):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    texts = json.load(open(os.path.join(HERE, "mail_texts.json"), encoding="utf-8"))
    tr = json.load(open(os.path.join(HERE, "xmb_translations.json"), encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    wrap = Alignment(wrap_text=True, vertical="top")
    for key, rows in texts["texts"].items():
        short = key.split("/")[-1].replace(".xmb", "")
        ws = wb.create_sheet(short[:28])
        ws.append(["파일", "id", "원문(JP)", "한국어(KR)", "메모"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            kr = (tr.get(key, {}).get(r["id"]) or {}).get("kr", "")
            ws.append([key, r["id"], r["jp"], kr, ""])
        for col, w in zip("ABCDE", (24, 12, 56, 56, 16)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "A2"
    if os.path.exists(XLSX_PATH):
        raise SystemExit(f"[중단] 이미 존재: {XLSX_PATH} — 기존 엑셀을 덮지 않음")
    wb.save(XLSX_PATH)
    n = sum(len(r) for r in texts["texts"].values())
    print(f"엑셀 생성: {XLSX_PATH} ({n}행, 시트 {len(texts['texts'])}개)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["dump", "export-xlsx"])
    args = ap.parse_args()
    {"dump": cmd_dump, "export-xlsx": cmd_export}[args.cmd](args)


if __name__ == "__main__":
    main()
