#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
구글 스프레드시트(팀이 쓰던 번역 시트) → 웹앱 script 분류 병합 (2026-10-23, 1단계: 드라이런)

  python sheet_merge.py dryrun <시트.xlsx>     # 서버 최신 rows·hist 기준 분류 보고(쓰기 없음)

시트 구조(시트마다 다름 — 머리글로 판별):
  인물명 | 원문 | 번역문(사람) | 비고 ... | Gemini | 파일경로 | 엔트리인덱스(=오프셋) | 원문바이트수
  all·ami·azu·ior·rit·ryu: 비고 4열(경로·오프셋·바이트 수가 비고 2~4열에 들어 있음) 뒤 Gemini
  yuk: Gemini 열 없음.  Gemini 값은 구글 수식(=IFERROR(__xludf...)) → 캐시값(data_only)으로 읽음.
행 대응: (시트 접두, "파일:오프셋") = 웹 script 키 → 원문이 같으면 확정. 안 맞으면 같은 파일 안에서
원문 정렬(difflib)로 재매칭(오프셋 체계 차이: 웹은 translation_2 오프셋 교정판).

분류(정책 초안 = KOR 권고: 웹 팀 편집 > 시트 사람 번역 > 기존 웹 기계번역, 시트 Gemini는 빈칸만):
  same            시트 사람 번역 = 웹 kr
  fill_human      웹 kr 빔 → 시트 사람 번역으로 채움
  overwrite       웹 kr(초기 Gemini API 기계번역) → 시트 사람 번역으로 교체
  conflict        웹에서 팀원이 편집한 행(by/hist) + 시트 사람 번역이 다름 → 덮지 않고 목록
  fill_gemini     사람 번역 없음 + 웹 kr 빔 → 시트 Gemini로 채움(출처 표시)
  gemini_same     시트 Gemini ≈ 웹 kr(문장부호·공백·개행 무시)
  gemini_alt      시트 Gemini ≠ 웹 kr → kr 유지, meta.alt 대안으로만
  unmatched       웹 행을 못 찾음
"""
import os
import re
import sys
import json
import difflib
import argparse
import collections
import urllib.request

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from extract_image_text import OUT_DEFAULT  # noqa: E402

CFG = json.load(open(os.path.join(HERE, "webapp_config.json"), encoding="utf-8"))
OUT = os.path.join(OUT_DEFAULT, "sheet_merge")      # 게임 원문 포함 — git 밖
SCB_RE = re.compile(r"^[a-z]{3}\\.+\.scb$", re.I)
PUNCT = re.compile(r"[\s「」『』（）()、。，．,.!！?？…~～・\-─—]+")


def norm_jp(s):
    return (s or "").replace("\r", "").strip().rstrip("\n").strip()


def loose(s):
    return PUNCT.sub("", s or "")


def read_sheet(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        if "(" not in ws.title:
            continue
        grp = ws.title.split("(", 1)[0]
        rows = ws.iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(rows)]
        c_jp, c_hu = head.index("원문"), head.index("번역문")
        c_gem = head.index("Gemini") if "Gemini" in head else None
        c_path = None
        for i, row in enumerate(ws.iter_rows(min_row=2, max_row=40, values_only=True)):
            for j, v in enumerate(row):
                if isinstance(v, str) and SCB_RE.match(v.strip()):
                    c_path = j
                    break
            if c_path is not None:
                break
        for n, row in enumerate(rows):
            row = list(row) + [None] * 30
            jp = row[c_jp]
            path = row[c_path]
            if jp is None and path is None:
                continue
            off = row[c_path + 1]
            gem = row[c_gem] if c_gem is not None else None
            if isinstance(gem, str) and (gem.startswith("=") or gem.startswith("#")):
                gem = None
            out.append({"grp": grp, "n": n, "sp": row[0], "jp": str(jp or ""), "hu": str(row[c_hu] or "").strip(),
                        "gem": str(gem or "").strip(), "file": str(path or "").strip(),
                        "off": int(float(off)) if off not in (None, "") else None})
    return out


def fetch(path_q):
    rows, off = [], 0
    while True:
        r = urllib.request.Request(f"{CFG['url']}/rest/v1/{path_q}&limit=1000&offset={off}", headers={
            "apikey": CFG["anon_key"], "Authorization": f"Bearer {CFG['anon_key']}", "x-team-code": "w@ldo"})
        ch = json.load(urllib.request.urlopen(r))
        rows += ch
        if len(ch) < 1000:
            return rows
        off += 1000


def match(sheet, web):
    """시트 행 → 웹 행. 1) 키+원문 일치 2) 파일별 원문 정렬 재매칭."""
    by_key = {}
    by_file = collections.defaultdict(list)
    for r in web:
        base = r["key"].split("#")[0]
        by_key.setdefault((r["grp"], base), r)
        by_file[(r["grp"], (r.get("meta") or {}).get("file", base.rsplit(":", 1)[0]))].append(r)
    pairs, used, how = {}, set(), {}
    rest = collections.defaultdict(list)
    for i, s in enumerate(sheet):
        w = by_key.get((s["grp"], f"{s['file']}:{s['off']}"))
        if w and norm_jp(w["jp"]) == norm_jp(s["jp"]) and w["id"] not in used:
            pairs[i], how[i] = w, "key"
            used.add(w["id"])
        else:
            rest[(s["grp"], s["file"])].append(i)
    for fk, idxs in rest.items():
        ws = [w for w in sorted(by_file.get(fk, []), key=lambda r: r["seq"]) if w["id"] not in used]
        a = [norm_jp(sheet[i]["jp"]) for i in idxs]
        b = [norm_jp(w["jp"]) for w in ws]
        sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
        for op, i1, i2, j1, j2 in sm.get_opcodes():
            if op == "equal":
                for k in range(i2 - i1):
                    pairs[idxs[i1 + k]], how[idxs[i1 + k]] = ws[j1 + k], "realign"
                    used.add(ws[j1 + k]["id"])
    return pairs, how


PH_ANY = re.compile(r"[\ue000-\uf8ff]")


def incomplete(jp, hu):
    """시트 '번역문'이 불완전한지(2026-10-23: 첫 줄만 번역된 행 1,768개 발견 → 되돌림).
    여러 줄 원문인데 한 줄뿐 / 원문은 닫는 괄호로 끝나는데 번역은 아님 / 치환 코드·이모티콘 불일치."""
    j, h = norm_jp(jp), hu.strip()
    w = []
    if "\n" in j and "\n" not in h:
        w.append("줄 누락")
    if j.endswith(("」", "）", "』")) and not h.endswith(("」", "）", "』", ")")):
        w.append("닫는 괄호 없음")
    if sorted(PH_ANY.findall(j)) != sorted(PH_ANY.findall(h)):
        w.append("치환코드·이모티콘 다름")
    return w


def classify(sheet, pairs, team):
    cls = collections.defaultdict(list)
    for i, s in enumerate(sheet):
        w = pairs.get(i)
        if w is None:
            if s["hu"] or s["gem"]:
                cls["unmatched"].append((s, None))
            continue
        kr = w["kr"] or ""
        if s["hu"]:
            if loose(kr) == loose(s["hu"]):
                cls["same"].append((s, w))
            elif incomplete(s["jp"], s["hu"]):
                cls["human_incomplete"].append((s, w))      # 덮지 않음(대안으로만)
            elif w["id"] in team:
                cls["conflict"].append((s, w))
            elif not kr:
                cls["fill_human"].append((s, w))
            else:
                cls["overwrite"].append((s, w))
        elif s["gem"]:
            if not kr:
                cls["fill_gemini"].append((s, w))
            elif loose(kr) == loose(s["gem"]):
                cls["gemini_same"].append((s, w))
            else:
                cls["gemini_alt"].append((s, w))
    return cls


def cmd_dryrun(a):
    sheet = read_sheet(a.xlsx)
    web = fetch("rows?select=id,grp,seq,key,jp,kr,st,by,ts,meta&cat=eq.script&order=id")
    hist = fetch("hist?select=row_id,by&cat=eq.script&order=row_id")
    team = {h["row_id"] for h in hist if h.get("by") and "초벌" not in h["by"]} | \
           {r["id"] for r in web if r.get("by") and "초벌" not in r["by"]}
    pairs, how = match(sheet, web)
    cls = classify(sheet, pairs, team)
    n_hu = sum(1 for s in sheet if s["hu"])
    n_gem = sum(1 for s in sheet if s["gem"])
    hu_rows = [s for s in sheet if s["hu"]]
    stat = {
        "sheet_rows": len(sheet), "human": n_hu, "gemini": n_gem,
        "web_script_rows": len(web), "web_team_edited": len(team),
        "matched_key": sum(1 for v in how.values() if v == "key"),
        "matched_realign": sum(1 for v in how.values() if v == "realign"),
        "classes": {k: len(v) for k, v in cls.items()},
        "human_by_sheet": dict(collections.Counter(s["grp"] for s in hu_rows)),
        "human_missing_newline": sum(1 for s in hu_rows if "\n" in norm_jp(s["jp"]) and "\n" not in s["hu"]),
        "human_missing_kagi": sum(1 for s in hu_rows if s["jp"].lstrip().startswith("「") and not s["hu"].lstrip().startswith("「")),
        "speaker_label_rows": sum(1 for s in sheet if re.fullmatch(r".{1,6}達", norm_jp(s["jp"]) or "")),
    }
    os.makedirs(OUT, exist_ok=True)
    ex = {k: [{"key": w["key"] if w else f"{s['file']}:{s['off']}", "grp": s["grp"], "jp": norm_jp(s["jp"])[:80],
               "web_kr": (w or {}).get("kr", "")[:80], "sheet_human": s["hu"][:80], "sheet_gemini": s["gem"][:80]}
              for s, w in v[:300]] for k, v in cls.items()}
    json.dump({"stat": stat, "examples": ex}, open(os.path.join(OUT, "dryrun.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(json.dumps(stat, ensure_ascii=False, indent=1))
    for k in ("fill_human", "overwrite", "conflict", "fill_gemini", "gemini_alt", "unmatched"):
        for s, w in cls.get(k, [])[:2]:
            print(f"[{k}] {s['grp']} {norm_jp(s['jp'])[:30]!r} | 웹 {(w or {}).get('kr', '')[:30]!r} | "
                  f"사람 {s['hu'][:30]!r} | Gemini {s['gem'][:30]!r}")
    print(f"→ {os.path.join(OUT, 'dryrun.json')}")


# ---------------------------------------------------------------- 반영 (sync, 멱등)
# 출처 meta.src_tr: human_sheet(구글시트 사람) / gemini_sheet(Gemini 초벌·시트) /
#                   machine_init(초기 기계번역 Gemini API, 웹 시딩) / claude(Claude 초벌) / team(웹 작업자)
BY_SHEET = "구글시트"
BY_GEM = "Gemini 초벌(시트)"
SYSTEM_BY = ("", BY_SHEET, BY_GEM, "Claude 초벌")
SHEET_URL = "https://docs.google.com/spreadsheets/d/1DDqH0g7Afpxf4hLVSBaWsaTnXiMd_v0baVaa5CEIpuk/export?format=xlsx"


def fix_kagi(jp, kr):
    """원문이 「…」인데 번역에 「」가 없으면 감싼다. (새 문자열, 보정했는지)"""
    j, k = norm_jp(jp), kr.strip()
    if j.startswith("「") and j.endswith("」") and not k.startswith("「") and not k.endswith("」"):
        return f"「{k}」", True
    return kr, False


def _req(method, path, body=None, prefer=None):
    h = {"apikey": CFG["anon_key"], "Authorization": f"Bearer {CFG['anon_key']}",
         "x-team-code": "w@ldo", "Content-Type": "application/json"}
    if prefer:
        h["Prefer"] = prefer
    data = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
    with urllib.request.urlopen(urllib.request.Request(CFG["url"] + path, data=data, method=method,
                                                       headers=h)) as r:
        t = r.read().decode("utf-8")
        return json.loads(t) if t else None


def cmd_sync(a):
    """시트 → 웹 1회 이관 겸 증분 동기화. 다시 돌려도 같은 결과(이미 반영된 행은 same).
    웹 작업자 편집 행(hist/by에 사람 이름)은 절대 덮지 않고 충돌 목록으로."""
    import time
    import urllib.parse
    src = a.xlsx
    if not src:
        src = os.path.join(OUT, "sheet_latest.xlsx")
        os.makedirs(OUT, exist_ok=True)
        open(src, "wb").write(urllib.request.urlopen(urllib.request.Request(
            SHEET_URL, headers={"User-Agent": "Mozilla/5.0"}), timeout=120).read())
    sheet = read_sheet(src)
    web = fetch("rows?select=id,cat,grp,seq,key,jp,kr,st,by,ts,meta&cat=eq.script&order=id")
    hist = fetch("hist?select=row_id,by&cat=eq.script&order=row_id")
    team = {h["row_id"] for h in hist if h.get("by") not in SYSTEM_BY and "초벌" not in (h.get("by") or "")
            and "기계번역" not in (h.get("by") or "")} | \
           {r["id"] for r in web if r.get("by") not in SYSTEM_BY}
    pairs, how = match(sheet, web)
    cls = classify(sheet, pairs, team)
    by_id = {r["id"]: r for r in web}
    sheet_of = {w["id"]: s for i, s in enumerate(sheet) if (w := pairs.get(i))}

    # 1) 메타: 출처·대안 (kr 열은 보내지 않음)
    meta_upd = []
    for r in web:
        m = dict(r.get("meta") or {})
        s = sheet_of.get(r["id"])
        if r["id"] in team:
            m["src_tr"] = "team"
        elif r.get("by") == BY_SHEET:
            m["src_tr"] = "human_sheet"
        elif r.get("by") == BY_GEM:
            m["src_tr"] = "gemini_sheet"
        elif r.get("by") == "Claude 초벌":
            m["src_tr"] = "claude"
        elif r["kr"]:
            m["src_tr"] = "machine_init"
        else:
            m.pop("src_tr", None)
        if s and s["gem"] and not s["hu"] and r["kr"] and loose(s["gem"]) != loose(r["kr"]):
            m["alt"] = {"kr": s["gem"], "src": "gemini_sheet"}
        if s and s["hu"] and incomplete(s["jp"], s["hu"]) and loose(s["hu"]) != loose(r["kr"] or ""):
            m["alt"] = {"kr": s["hu"], "src": "human_sheet",
                        "note": "시트 번역(" + "·".join(incomplete(s["jp"], s["hu"])) + " — 불완전해 적용 안 함)"}
            m["sheet_rejected"] = "·".join(incomplete(s["jp"], s["hu"]))
        if m != (r.get("meta") or {}):
            meta_upd.append({k: r[k] for k in ("id", "cat", "grp", "seq", "key", "jp")} | {"meta": m})
            r["meta"] = m
    print(f"메타(출처·대안) 갱신 {len(meta_upd)}행" + (" (드라이런)" if a.dry_run else ""))
    if not a.dry_run:
        for i in range(0, len(meta_upd), 500):
            _req("POST", "/rest/v1/rows?on_conflict=id", meta_upd[i:i + 500],
                 "resolution=merge-duplicates,return=minimal")

    # 2) kr 반영: 사람 번역(교체·채움), Gemini는 빈칸만
    jobs = []
    for kind in ("overwrite", "fill_human", "fill_gemini"):
        for s, w in cls.get(kind, []):
            if kind == "fill_gemini":
                new, fixed, by, srct = s["gem"], False, BY_GEM, "gemini_sheet"
            else:
                new, fixed = fix_kagi(s["jp"], s["hu"])
                by, srct = BY_SHEET, "human_sheet"
            jobs.append((kind, w, new, fixed, by, srct))
    done = collections.Counter()
    hist_rows = []
    for kind, w, new, fixed, by, srct in jobs:
        r = by_id[w["id"]]
        m = dict(r.get("meta") or {})
        m["src_tr"] = srct
        if r["kr"] and "prev_kr" not in m:
            m["prev_kr"] = r["kr"]
            m["prev_src"] = "machine_init"
        if fixed:
            m["merge_note"] = "「」 자동 보정(시트 번역에 없음)"
        done[kind] += 1
        done["kagi"] += fixed
        if a.dry_run:
            continue
        ts = int(time.time() * 1000)
        q = f"/rest/v1/rows?id=eq.{urllib.parse.quote(r['id'], safe='')}&kr=eq.{urllib.parse.quote(r['kr'] or '', safe='')}"
        res = _req("PATCH", q, {"kr": new, "st": max(1, r["st"]), "ts": ts, "by": by, "meta": m},
                   "return=representation")
        if not res:
            done["skipped_race"] += 1
            continue
        if r["kr"]:
            hist_rows.append({"row_id": r["id"], "cat": "script", "grp": r["grp"], "key": r["key"],
                              "ts": ts - 1, "by": "기계번역(교체 전)", "kr": r["kr"], "st": r["st"]})
        hist_rows.append({"row_id": r["id"], "cat": "script", "grp": r["grp"], "key": r["key"],
                          "ts": ts, "by": by, "kr": new, "st": max(1, r["st"])})
        print(f"  kr 반영 {sum(done[k] for k in ('overwrite', 'fill_human', 'fill_gemini'))}/{len(jobs)}", end="\r")
    for i in range(0, len(hist_rows), 200):
        _req("POST", "/rest/v1/hist", hist_rows[i:i + 200], "return=minimal")

    # 3) 목록 파일
    os.makedirs(OUT, exist_ok=True)
    conf = [{"id": w["id"], "jp": w["jp"], "web_kr": w["kr"], "web_by": w.get("by"), "sheet_human": s["hu"]}
            for s, w in cls.get("conflict", [])]
    unm = [{"grp": s["grp"], "file": s["file"], "off": s["off"], "jp": s["jp"], "sheet_human": s["hu"],
            "sheet_gemini": s["gem"]} for s, _w in cls.get("unmatched", [])]
    realign = [{"id": pairs[i]["id"], "sheet_key": f"{sheet[i]['file']}:{sheet[i]['off']}", "jp": sheet[i]["jp"]}
               for i, h in how.items() if h == "realign"]
    for fn, data in (("conflicts.json", conf), ("unmatched_manual.json", unm), ("realigned.json", realign)):
        json.dump(data, open(os.path.join(OUT, fn), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    rep = {"at": time.strftime("%Y-%m-%d %H:%M"), "dry_run": a.dry_run, "overwrite": done["overwrite"],
           "fill_human": done["fill_human"], "fill_gemini": done["fill_gemini"], "kagi_fixed": done["kagi"],
           "skipped_race": done["skipped_race"], "alt": sum(1 for r in web if (r.get("meta") or {}).get("alt")),
           "same": len(cls.get("same", [])), "conflict": len(conf), "unmatched": len(unm), "realigned": len(realign),
           "hist": len(hist_rows)}
    json.dump(rep, open(os.path.join(OUT, "sync_report.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("\n" + json.dumps(rep, ensure_ascii=False))


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("dryrun")
    d.add_argument("xlsx")
    s = sub.add_parser("sync", help="시트 → 웹 반영(멱등). xlsx 생략 시 공유 링크에서 내려받음")
    s.add_argument("xlsx", nargs="?")
    s.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    {"dryrun": cmd_dryrun, "sync": cmd_sync}[a.cmd](a)


if __name__ == "__main__":
    main()
