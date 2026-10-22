#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
iM@S2 KOR 번역 플랫폼 회수 (Supabase rows → 파이프라인 입력)

  python webapp_pull.py xmb    --code <암호> [--reviewed-only]
      → xmb_translations.json 병합 (이후 xmb_text.py build)
  python webapp_pull.py script --code <암호> [--reviewed-only]
      → 게임 tools에 translation_webapp_<날짜>.xlsx 생성
        (translation_2.xlsx와 같은 7열 포맷, 기존 파일 안 덮음.
         excel2scb 전환은 scb-excel 쪽에서 파일명 교체로)
  python webapp_pull.py image  --code <암호> [--reviewed-only]
      → label_specs.json의 items[md5].text 병합 (이후 label_tool --export)
"""
import os
import sys
import json
import time
import argparse
import urllib.request

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from extract_image_text import JP_ROOT  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
CFG = json.load(open(os.path.join(HERE, "webapp_config.json"), encoding="utf-8"))
GRP2REL = {"install": "text/install.xmb", "ranker": "ranker/ranker_npc.xmb",
           "exep_init": "scene/exep_init.mpc"}


def fetch(cat, code):
    rows = []
    offset = 0
    while True:
        url = (f"{CFG['url']}/rest/v1/rows?cat=eq.{cat}&select=*"
               f"&order=grp,seq&limit=1000&offset={offset}")
        r = urllib.request.Request(url, headers={
            "apikey": CFG["anon_key"], "Authorization": f"Bearer {CFG['anon_key']}",
            "x-team-code": code})
        chunk = json.load(urllib.request.urlopen(r))
        rows += chunk
        if len(chunk) < 1000:
            break
        offset += 1000
    print(f"{cat}: 서버 {len(rows)}행")
    return rows


def want(r, reviewed_only):
    return bool(r.get("kr")) and (not reviewed_only or r.get("st", 0) >= 2)


def pull_xmb(rows, reviewed_only):
    tr_path = os.path.join(HERE, "xmb_translations.json")
    tr = json.load(open(tr_path, encoding="utf-8"))
    n = 0
    for r in rows:
        if not want(r, reviewed_only):
            continue
        if r["grp"] == "lyrics":            # key = 곡코드/LYR_nnn
            song, _, tid = r["key"].partition("/")
            rel = f"songinfo/{song}/lyrics_{song}.xmb"
        else:
            rel = GRP2REL.get(r["grp"])
            tid = r["key"]
        if not rel:
            continue
        tid = tid.split("#")[0]
        if tid.startswith("x360:"):         # 360판 NPC 명단에서 원문이 다른 칸 (pcport_apply.npc_kr_for)
            rel, tid = rel + "@x360", tid[5:]
        ent = tr.setdefault(rel, {}).setdefault(tid, {"jp": r["jp"], "cn": "", "kr": ""})
        if ent.get("kr") != r["kr"]:
            ent["kr"] = r["kr"]
            n += 1
    json.dump(tr, open(tr_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"xmb_translations.json 반영 {n}건 → 다음: python xmb_text.py build")


def pull_script(rows, reviewed_only):
    """translation_2.xlsx를 베이스로 kr만 갱신한 새 xlsx 스냅샷 생성."""
    from openpyxl import load_workbook
    src = os.path.join(JP_ROOT, "gscript", "tools", "translation_2.xlsx")
    dst = os.path.join(JP_ROOT, "gscript", "tools",
                       f"translation_webapp_{time.strftime('%Y%m%d_%H%M')}.xlsx")
    upd = {}
    for r in rows:
        if want(r, reviewed_only):
            upd[(r["grp"], r["key"].split("#")[0])] = r["kr"]
    wb = load_workbook(src)
    n = 0
    for ws in wb.worksheets:
        if "(" not in ws.title:
            continue
        grp = ws.title.split("(", 1)[0]
        for row in ws.iter_rows(min_row=2):
            file, off = row[4].value, row[5].value
            if file is None:
                continue
            kr = upd.get((grp, f"{file}:{off}"))
            if kr is not None and row[2].value != kr:
                row[2].value = kr
                n += 1
    wb.save(dst)
    print(f"갱신 {n}셀 → {dst}")
    # PC 포트(360) 전용 대사 → translation_x360_<날짜>.json (pcport360.py build가 최신 파일 사용)
    x = {}
    for r in rows:
        m = r.get("meta") or {}
        if m.get("x360") and want(r, reviewed_only):
            x.setdefault(m["file"], {})[str(m["x_idx"])] = {"jp": r["jp"], "kr": r["kr"]}
    if x:
        xp = os.path.join(JP_ROOT, "gscript", "tools",
                          f"translation_x360_{time.strftime('%Y%m%d_%H%M')}.json")
        json.dump(x, open(xp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print(f"360 전용 대사 {sum(len(v) for v in x.values())}줄 → {xp}")
    print("excel2scb 전환은 scb-excel 파이프라인에서 이 파일로 교체 후 실행")


def pull_image(rows, reviewed_only, code=None):
    import base64
    from extract_image_text import OUT_DEFAULT
    sp_path = os.path.join(HERE, "label_specs.json")
    specs = json.load(open(sp_path, encoding="utf-8")) if os.path.exists(sp_path) \
        else {"presets": {}, "items": {}}
    n = 0
    for r in rows:
        if not want(r, reviewed_only):
            continue
        it = specs["items"].setdefault(r["key"], {"preset": "기본", "over": {},
                                                  "mode": "text"})
        if it.get("text") != r["kr"]:
            it["text"] = r["kr"]
            n += 1
    json.dump(specs, open(sp_path, "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(f"label_specs.json 문구 반영 {n}건")

    # 웹앱 업로드 작업본(tex_kr) → kr/ PNG (md5의 모든 위치에 복제)
    by_key = {r["key"]: r for r in rows}
    imgs = []
    offset = 0
    while True:
        req_ = urllib.request.Request(
            f"{CFG['url']}/rest/v1/tex_kr?select=*&limit=100&offset={offset}",
            headers={"apikey": CFG["anon_key"],
                     "Authorization": f"Bearer {CFG['anon_key']}",
                     "x-team-code": code})
        chunk = json.load(urllib.request.urlopen(req_))
        imgs += chunk
        if len(chunk) < 100:
            break
        offset += 100
    m = 0
    for im in imgs:
        r = by_key.get(im["md5"])
        if r is None:
            continue
        if reviewed_only and r.get("st", 0) < 2:
            continue
        png = base64.b64decode(im["data"])
        for loc in (r.get("meta") or {}).get("locs", []):
            rel, nn = loc.rsplit("#", 1)
            if loc.startswith("x360:"):  # 360 본편 전용 그림: kr_x360/<상대경로>_<n>.png
                stem = rel[5:].replace("/", "_").rsplit(".", 1)[0]
                dst = os.path.join(OUT_DEFAULT, "kr_x360", f"{stem}_{nn}.png")
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                with open(dst, "wb") as f:
                    f.write(png)
                m += 1
                continue
            if loc.startswith("360:"):   # DLC: kr_dlc/<패키지6>/<상대경로>_<n>.png
                pkg, path = rel[4:].split("/", 1)
                stem = path.replace("/", "_").rsplit(".", 1)[0]
                dst = os.path.join(OUT_DEFAULT, "kr_dlc", pkg, f"{stem}_{nn}.png")
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                with open(dst, "wb") as f:
                    f.write(png)
                m += 1
                continue
            rel = rel.replace("/", "\\")
            stem = os.path.splitext(os.path.basename(rel))[0]
            dst = os.path.join(OUT_DEFAULT, "kr", os.path.dirname(rel),
                               f"{stem}_{nn}.png")
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            with open(dst, "wb") as f:
                f.write(png)
            m += 1
    print(f"작업본 PNG {m}장 저장 ({len(imgs)}개 업로드본, 모든 위치 복제)")
    print("다음: python label_tool.py --export (라벨 렌더) → python inject_image_text.py")


def pull_dlc(rows, reviewed_only):
    """dlc 행 회수 → dlc_translations.json (grp/key→kr, meta.src 포함).
    songinfo 그룹은 xmb_translations.json의 songinfo/<곡> 키에도 병합해
    dance_lyrics 계열 빌드가 재사용할 수 있게 한다 (DLC dance pac 주입은 별도)."""
    out = {}
    n = 0
    for r in rows:
        if not want(r, reviewed_only):
            continue
        out[f"{r['grp']}/{r['key']}"] = {"jp": r["jp"], "kr": r["kr"],
                                         "meta": r.get("meta", {})}
        n += 1
    p = os.path.join(HERE, "dlc_translations.json")
    json.dump(out, open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"dlc_translations.json {n}건")
    # songinfo → xmb_translations 병합 (lyrics_<곡>:LYR_nnn → songinfo/<곡>)
    tr_path = os.path.join(HERE, "xmb_translations.json")
    tr = json.load(open(tr_path, encoding="utf-8"))
    m = 0
    for r in rows:
        if r["grp"] != "songinfo" or not want(r, reviewed_only):
            continue
        stem, tid = r["key"].split(":", 1)
        song = stem.replace("lyrics_", "")
        ent = tr.setdefault(f"songinfo/{song}/lyrics_{song}.xmb", {}) \
            .setdefault(tid, {"jp": r["jp"], "cn": "", "kr": ""})
        if ent.get("kr") != r["kr"]:
            ent["kr"] = r["kr"]
            m += 1
    json.dump(tr, open(tr_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"xmb_translations.json songinfo 병합 {m}건 (DLC 신곡 가사)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cat", choices=["xmb", "script", "image", "dlc"])
    ap.add_argument("--code", required=True)
    ap.add_argument("--reviewed-only", action="store_true")
    args = ap.parse_args()
    rows = fetch(args.cat, args.code)
    if args.cat == "image":
        pull_image(rows, args.reviewed_only, code=args.code)
    else:
        {"xmb": pull_xmb, "script": pull_script,
         "dlc": pull_dlc}[args.cat](rows, args.reviewed_only)


if __name__ == "__main__":
    main()
