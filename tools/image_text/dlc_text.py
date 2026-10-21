#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DLC 번역 대상 추출 (360 기준 축 + PS3 병합) — imas2.buyeo.uk 'dlc' 분류용

소스 (읽기 전용, 사본만 복호화):
  360 (PC 포트 주입 대상, 1순위):
    Desktop\\imas2dlc\\content\\4E4D0817\\00000002\\<pkg>\\resource\\
      catalog_info.xmb(암호화) / packageInfo.pac(평문 TSK) /
      mail\\*_mail_01.xmb + all_mtbl_01.xmb(1개 패키지만, 암호화) /
      songinfo\\<곡>\\lyrics_<곡>.xmb(암호화, DLC 신곡 가사)
  PS3 (업로드 zip 추출본, 보조·병합):
    dlc\\cpNN\\resource\\ packageInfo.pac(암호화) / mail\\*(cp01) / ranker\\npc_*.xmb

행 모델 (cat="dlc"):
  grp: mail / catalog / package / songinfo / ranker
  key: <파일스템>:<id>  (양 플랫폼 jp 동일 → 1행, meta.src에 두 위치 기록.
       jp가 다르면 "#ps3" 접미 별도 행)
  kr 제안: 본편 기번역(jp 전문 일치)을 xmb_translations.json + translation_2.xlsx에서

사용:
  python dlc_text.py extract [--ps3-root <폴더>]   # dlc_texts.json + 드라이런 보고
  (시딩은 webapp_seed.py --only dlc --code <팀암호> — 원격 쓰기는 사용자 실행)
"""
import os
import re
import sys
import json
import shutil
import argparse
import subprocess
from collections import Counter, defaultdict

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from xmb_text import parse_xmb, collect_texts, MAPPING_PATH  # noqa: E402
from extract_image_text import OUT_DEFAULT, IMAS2DEC, JP_ROOT  # noqa: E402

X360_ROOT = r"C:\Users\jackfilm\Desktop\imas2dlc\content\4E4D0817\00000002"
PS3_ROOT_DEFAULT = os.path.join(
    os.environ.get("TEMP", ""), "")  # 실행 시 --ps3-root로 지정 (zip 추출 위치)
CACHE = os.path.join(OUT_DEFAULT, "_cache", "dlc")
PLAIN = (b"XMB ", b"TSK ", b"pac\x00")


def stage(src: str, tag: str) -> bytes:
    """사본을 캐시에 뜨고 암호화면 복호화. 원본 불변."""
    dst = os.path.join(CACHE, tag.replace("/", "_").replace("\\", "_"))
    if not os.path.exists(dst):
        os.makedirs(CACHE, exist_ok=True)
        shutil.copy2(src, dst)
        if open(dst, "rb").read(4) not in PLAIN:
            p = subprocess.Popen([IMAS2DEC, dst], cwd=CACHE,
                                 stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            try:
                p.wait(120)
            except subprocess.TimeoutExpired:
                p.kill()
            if not os.path.exists(dst + ".dec"):
                raise RuntimeError(f"복호화 실패: {src}")
            os.replace(dst + ".dec", dst)
    return open(dst, "rb").read()


def xmb_texts_all(data: bytes):
    """버퍼의 모든 XMB에서 (xmb_idx, id, text). 파일 내 중복 id는 '.n' 서수
    (같은 아이템 노드의 이름/설명 text_node가 조상 id를 공유하는 구조)."""
    out = []
    seen = Counter()
    p = 0
    xi = 0
    while True:
        i = data.find(b"XMB ", p)
        if i < 0:
            break
        try:
            ts = collect_texts(parse_xmb(data, i))
            for t in ts:
                k = (xi, t["id"])
                seen[k] += 1
                tid = t["id"] if seen[k] == 1 else f"{t['id']}.{seen[k]}"
                out.append((xi, tid, t["text"]))
        except Exception:  # noqa: BLE001 — 깨진 후보 스킵
            pass
        xi += 1
        p = i + 4
    return out


# ---------------------------------------------------------------- 수집

def collect(ps3_root):
    """rows[(grp, stem, id)] = {jp, src:[...]} — 360 먼저(기준), PS3 병합."""
    rows = {}
    clashes = []

    def add(grp, stem, tid, jp, src):
        if not jp or jp.isascii():       # 비ASCII 텍스트만 번역 대상
            return
        k = (grp, stem, tid)
        if k not in rows:
            rows[k] = {"jp": jp, "src": [src]}
        elif rows[k]["jp"] == jp:
            rows[k]["src"].append(src)
        else:
            k2 = (grp, stem, tid + "#ps3")
            if k2 not in rows:
                rows[k2] = {"jp": jp, "src": [src]}
                clashes.append(k)
            else:
                rows[k2]["src"].append(src)

    # ---- 360 (기준)
    pkgs = sorted(os.listdir(X360_ROOT))
    for pkg in pkgs:
        res = os.path.join(X360_ROOT, pkg, "resource")
        if not os.path.isdir(res):
            continue
        tagp = f"360_{pkg[:6]}"
        ci = os.path.join(res, "catalog_info.xmb")
        if os.path.exists(ci):
            # catalog_info는 XMB 여러 개(이름/설명 등)에 같은 id가 반복 → 서수로 구분
            for xi, tid, tx in xmb_texts_all(stage(ci, f"{tagp}_catalog_info.xmb")):
                add("catalog", f"catalog_{pkg[:6]}", f"x{xi}.{tid}", tx,
                    f"360:{pkg[:6]}/catalog_info.xmb#x{xi}")
        pi = os.path.join(res, "packageInfo.pac")
        if os.path.exists(pi):
            for xi, tid, tx in xmb_texts_all(stage(pi, f"{tagp}_packageInfo.pac")):
                add("package", f"package_{pkg[:6]}", f"x{xi}.{tid}", tx,
                    f"360:{pkg[:6]}/packageInfo.pac#x{xi}")
        md = os.path.join(res, "mail")
        if os.path.isdir(md):
            for fn in sorted(os.listdir(md)):
                if fn.endswith(".xmb"):
                    stem = fn[:-4]
                    for _xi, tid, tx in xmb_texts_all(
                            stage(os.path.join(md, fn), f"{tagp}_{fn}")):
                        add("mail", stem, tid, tx, f"360:{pkg[:6]}/mail/{fn}")
        si = os.path.join(res, "songinfo")
        if os.path.isdir(si):
            for song in sorted(os.listdir(si)):
                ly = os.path.join(si, song, f"lyrics_{song}.xmb")
                if os.path.exists(ly):
                    for _xi, tid, tx in xmb_texts_all(
                            stage(ly, f"{tagp}_lyrics_{song}.xmb")):
                        add("songinfo", f"lyrics_{song}", tid, tx,
                            f"360:{pkg[:6]}/songinfo/{song}")

    # ---- PS3 (병합)
    dlc = os.path.join(ps3_root, "BLJS10083", "USRDIR", "dlc")
    for cp in sorted(os.listdir(dlc)):
        res = os.path.join(dlc, cp, "resource")
        if not os.path.isdir(res):
            continue
        pi = os.path.join(res, "packageInfo.pac")
        if os.path.exists(pi):
            # PS3 packageInfo는 360과 패키지 대응이 불명 → jp 전문 일치로 병합
            jp_index = {v["jp"]: k for k, v in rows.items() if k[0] == "package"}
            for xi, tid, tx in xmb_texts_all(stage(pi, f"ps3_{cp}_packageInfo.pac")):
                if not tx or tx.isascii():
                    continue
                hit = jp_index.get(tx)
                if hit:
                    rows[hit]["src"].append(f"ps3:{cp}/packageInfo.pac#x{xi}")
                else:
                    add("package", f"package_{cp}", f"x{xi}.{tid}", tx,
                        f"ps3:{cp}/packageInfo.pac#x{xi}")
        md = os.path.join(res, "mail")
        if os.path.isdir(md):
            for fn in sorted(os.listdir(md)):
                if fn.endswith(".xmb"):
                    for _xi, tid, tx in xmb_texts_all(
                            stage(os.path.join(md, fn), f"ps3_{cp}_{fn}")):
                        add("mail", fn[:-4], tid, tx, f"ps3:{cp}/mail/{fn}")
        rd = os.path.join(res, "ranker")
        if os.path.isdir(rd):
            for fn in sorted(os.listdir(rd)):
                if fn.endswith(".xmb"):
                    for _xi, tid, tx in xmb_texts_all(
                            stage(os.path.join(rd, fn), f"ps3_{cp}_{fn}")):
                        add("ranker", fn[:-4], tid, tx, f"ps3:{cp}/ranker/{fn}")
    return rows, clashes


# ---------------------------------------------------------------- 기번역 제안

def main_game_krmap():
    """본편 jp→kr (전문 일치): xmb_translations.json + translation_2.xlsx."""
    m = defaultdict(Counter)
    tr = json.load(open(os.path.join(HERE, "xmb_translations.json"), encoding="utf-8"))
    for ents in tr.values():
        for v in ents.values():
            if v.get("jp") and v.get("kr"):
                m[v["jp"]][v["kr"]] += 1
    from openpyxl import load_workbook
    xp = os.path.join(JP_ROOT, "gscript", "tools", "translation_2.xlsx")
    wb = load_workbook(xp, read_only=True)
    for ws in wb.worksheets:
        if "(" not in ws.title:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            jp, kr = row[1], row[2]
            if jp and kr:
                m[str(jp)][str(kr)] += 1
    return {jp: c.most_common(1)[0][0] for jp, c in m.items()}


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["extract"])
    ap.add_argument("--ps3-root", required=True,
                    help="PS3 DLC zip을 풀어둔 폴더 (BLJS10083 상위)")
    args = ap.parse_args()

    rows, clashes = collect(args.ps3_root)
    krmap = main_game_krmap()
    mapping = set(json.load(open(MAPPING_PATH, encoding="utf-8")).keys())

    out = {}
    seqs = Counter()
    n_pre = 0
    miss_chars = set()
    for (grp, stem, tid), v in sorted(rows.items()):
        key = f"{stem}:{tid}"
        kr = krmap.get(v["jp"], "")
        if kr:
            n_pre += 1
            miss_chars |= {c for c in kr if "가" <= c <= "힣" and c not in mapping}
        seqs[grp] += 1
        out[f"dlc:{grp}:{key}"] = {
            "cat": "dlc", "grp": grp, "seq": seqs[grp], "key": key,
            "jp": v["jp"], "kr": kr, "st": 1 if kr else 0,
            "meta": {"src": v["src"]}}
    with open(os.path.join(HERE, "dlc_texts.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)

    # ---- 드라이런 보고
    print(f"\n== DLC 추출 드라이런 — 총 {len(out)}행")
    for grp, n in sorted(seqs.items()):
        both = sum(1 for v in out.values()
                   if v["grp"] == grp
                   and any(s.startswith("360") for s in v["meta"]["src"])
                   and any(s.startswith("ps3") for s in v["meta"]["src"]))
        p360 = sum(1 for v in out.values() if v["grp"] == grp
                   and all(s.startswith("360") for s in v["meta"]["src"]))
        pps3 = sum(1 for v in out.values() if v["grp"] == grp
                   and all(s.startswith("ps3") for s in v["meta"]["src"]))
        pre = sum(1 for v in out.values() if v["grp"] == grp and v["kr"])
        print(f"  {grp:9s}: {n:5d}행 (양플랫폼 {both} / 360전용 {p360} / PS3전용 {pps3}"
              f" / 기번역 제안 {pre})")
    print(f"  jp 불일치로 분리된 행(#ps3): {len(clashes)}")
    print(f"  기번역 제안 합계: {n_pre}")
    print(f"  제안 kr 중 폰트 매핑(2,359자) 미보유 한글: "
          f"{len(miss_chars)}자 {''.join(sorted(miss_chars)) or '없음'}")
    print("\n== 샘플 5행:")
    for k, v in list(out.items())[:300:60]:
        print(f"  [{v['grp']}] {v['key']}")
        print(f"    JP: {v['jp'][:60].replace(chr(10), '⏎')}")
        if v["kr"]:
            print(f"    KR(제안): {v['kr'][:60].replace(chr(10), '⏎')}")
    print("\n시딩(사용자 실행): python webapp_seed.py --only dlc --code <팀암호>")


if __name__ == "__main__":
    main()
